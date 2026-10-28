"""PageIndex's LocalAPI with Textract + VLM for scanned PDFs, figures and tables.

PageIndex itself is unchanged and imported, not copied. Every submit runs through
upstream's ``LocalAPI.submit_document`` and ``page_index_flash``; this module plugs
into them at three points:

* ``_extract_page_texts``: the pages with a large visual region (a scanned page, a
  photo, a chart, a ruled table), and only those, are parsed by Textract. Their text
  (tables as markdown, figures described by the VLM from their S3 URL, headers and
  footers kept) replaces those pages' text the agent reads; every other page keeps
  its original PyPDF2 text. A digital PDF falls back to the original text on every
  page when Textract, S3 or the VLM fails. Upstream calls it after validating the
  submit's arguments, so a refused submit is never billed.
* ``extract_toc`` (patched in ``pageindex.flash.api`` by ``install``): Flash's own
  tree, with Textract's text in ``page_texts`` on the pages Textract read, so
  summaries and expand see the tables and figure descriptions. A PDF Flash cannot
  read (no text layer, no bookmarks) gets its tree from the Textract headings.
* ``_unique_doc_name``: every submit is indexed from a ``<uuid>.pdf`` in a work dir
  and stored under the submitted file's own name.

A ``.jpg`` / ``.jpeg`` / ``.png`` is indexed as a one-page PDF, and an office file
(``office_to_pdf.SUPPORTED_EXTENSIONS``) as the PDF unoserver renders from it. An
image Textract reads as mostly picture (a photo, a chart) goes to the VLM whole, its
OCR words kept beside the description; a document image (a scan, a screenshot)
stays OCR. A ``.xlsx`` is indexed like a PDF whose pages are its sheets: a blank
page per sheet, one tree node per sheet, each page's text opened by ``Sheet: <name>``.
"""
import asyncio
import copy
import logging
import os
import shutil
import tempfile
import threading
import uuid
from contextvars import ContextVar
from dataclasses import dataclass, field
from pathlib import Path

from pageindex.errors import PageIndexAPIError
from pageindex.local_api import LocalAPI
from pageindex.naming import sanitize_filename
from pageindex.utils import run_off_loop

logger = logging.getLogger(__name__)

IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
SPREADSHEET_EXTS = {".xlsx"}
HEADER_TYPES = {"LAYOUT_HEADER"}
FOOTER_TYPES = {"LAYOUT_FOOTER", "LAYOUT_PAGE_NUMBER"}
HEADING_MIN_CHARS, HEADING_MAX_CHARS = 3, 200
# Opens each sheet's page text; citation.py reads the sheet name back from it.
SHEET_PREFIX = "Sheet: "
# Set to a directory to keep Textract's raw/formatted JSON and the figure crops.
DEBUG_DIR_ENV = "PAGEINDEX_MM_DEBUG_DIR"


@dataclass
class Parsed:
    content_list: list
    markdown: list   # one string per page -> pages.json
    plain: list      # one string per page -> Flash's page_texts (no '#', no header/footer)
    visual_pages: list = field(default_factory=list)  # 0-based pages Textract read
    structure: list | None = None  # a spreadsheet's tree: one node per sheet


# The Parsed of the document being indexed. Set by MultimodalLocalAPI._index_flash on
# the thread that then runs page_index_flash, read by the extract_toc hook.
_CURRENT: ContextVar[Parsed | None] = ContextVar("pageindex_mm_parsed", default=None)


# ── pure helpers ──

def image_to_pdf(src: Path, dst: Path) -> Path:
    """One-page PDF at the render DPI, so the 400-DPI render is the image's own pixels."""
    from PIL import Image, ImageOps
    from .parser.figures import _flatten
    from .parser.pdf.amazon_textract.pages import RENDER_DPI
    with Image.open(src) as img:
        # EXIF: a phone photo would reach Textract and the VLM sideways
        _flatten(ImageOps.exif_transpose(img)).save(dst, "PDF", resolution=RENDER_DPI)
    return dst


def has_text_layer(pdf_path) -> bool:
    import pymupdf
    with pymupdf.open(pdf_path) as doc:
        return any(page.get_text().strip() for page in doc)


def textract_pages(pdf_path) -> list[int]:
    """0-based pages with a visual region Textract would keep (same ``is_large_enough``).

    Only decides which pages to pay Textract for (it bills per page, so a page is
    the unit); Textract decides what a figure or a table is. Raster images cover
    scans, photos and screenshots; drawing clusters cover vector charts and ruled
    tables. Every other page keeps its original PyPDF2 text.
    ponytail: a table without rules in a digital PDF is missed and keeps PyPDF2 text;
    send every page to Textract if that matters.
    """
    import pymupdf
    from .parser.figures import is_large_enough
    from .parser.pdf.amazon_textract.pages import RENDER_DPI
    scale = RENDER_DPI / 72
    pages = []
    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            page_area = page.rect.width * page.rect.height
            if not page_area:
                continue
            regions = [pymupdf.Rect(info["bbox"]) for info in page.get_image_info()]
            regions += page.cluster_drawings()
            for rect in regions:
                rect = rect & page.rect
                if not rect.is_empty and is_large_enough(
                        rect.width * scale, rect.height * scale,
                        rect.width * rect.height / page_area):
                    pages.append(page.number)
                    break
    return pages


def remap_pages(content_list: list, raw_parse: dict, pages: list[int]) -> None:
    """Point a sub-PDF's ``page_idx`` (0..k-1) back at the original pages, in place."""
    for item in content_list:
        item["page_idx"] = pages[item["page_idx"]]
    for raw_page in raw_parse.get("pages", []):
        raw_page["page_idx"] = pages[raw_page["page_idx"]]


def merge_pages(textract_pages_text: list[str], original: list[str], pages: list[int]) -> list[str]:
    """Textract's text on the pages it read, the original text everywhere else."""
    chosen = set(pages)
    return [textract_pages_text[i] if i in chosen else original[i] for i in range(len(original))]


# ── image files: document or visual ──

# Below this share of the image read as text or table, an image is visual (a photo, a
# chart): Textract would crop a part of it as a figure and drop the rest. Measured on
# Textract layout: document images (scans, screenshots) 34-57%, photos and charts 0-3%.
VISUAL_TEXT_COVERAGE = 0.15


def text_coverage(raw_page: dict) -> float:
    """Share of a page's area Textract lays out as text or table (header/footer included)."""
    from .parser.pdf.amazon_textract.blocks import build_block_map, get_children
    from .parser.pdf.amazon_textract.content_list import IGNORED_LAYOUT_TYPES, LAYOUT_TEXT_TYPES
    textual = LAYOUT_TEXT_TYPES | IGNORED_LAYOUT_TYPES | {"LAYOUT_TABLE"}
    blocks = raw_page.get("Blocks") or []
    page = next((b for b in blocks if b["BlockType"] == "PAGE"), None)
    if page is None:
        return 0.0
    boxes = [b["Geometry"]["BoundingBox"] for b in get_children(page, build_block_map(blocks))
             if b["BlockType"] in textual]
    return sum(box["Width"] * box["Height"] for box in boxes)


def as_whole_image(content_list: list, image_path: Path) -> list:
    """A visual image: the whole picture for the VLM, followed by the words Textract
    read on it, kept verbatim and handed to the VLM as its context."""
    whole = {"type": "image", "img_path": str(image_path), "image_caption": None,
             "image_footnote": None, "page_idx": 0}
    return [whole, *(item for item in content_list if item.get("type") == "text")]


def image_for_vlm(src: Path, tmp: Path) -> Path:
    """The submitted image, upright and within the VLM limits, as a file to upload."""
    from PIL import Image, ImageOps
    from .parser.figures import _flatten, save_for_vlm
    with Image.open(src) as img:
        return save_for_vlm(_flatten(ImageOps.exif_transpose(img)), tmp / "whole_image.jpg")


def _page_layout_texts(raw_page: dict, types: set) -> list:
    """Text of the header / footer layout blocks Textract found on one page."""
    from .parser.pdf.amazon_textract.blocks import build_block_map, get_children, get_plain_text
    blocks = raw_page.get("Blocks") or []
    if raw_page.get("error") or not blocks:
        return []
    block_map = build_block_map(blocks)
    page_block = next((b for b in blocks if b["BlockType"] == "PAGE"), None)
    if page_block is None:
        return []
    texts = (get_plain_text(b, block_map) for b in get_children(page_block, block_map)
             if b["BlockType"] in types)
    return [t for t in texts if t]


def _norm(text: str) -> str:
    return " ".join(text.split()).casefold()


def _item_block(item: dict, recent_text: str = "") -> tuple[str, str] | None:
    """(markdown, plain) of one content-list item, or None when it carries nothing.

    A table caption already written by the text just above it (Textract often reports
    the same line as both a layout text and the table title) is not repeated.
    """
    kind = item.get("type")
    if kind == "text":
        text = (item.get("text") or "").strip()
        if not text:
            return None
        level = item.get("text_level")
        return (f"{'#' * level} {text}" if level else text), text
    if kind == "table":
        captions = [c for c in item.get("table_caption") or [] if _norm(c) not in recent_text]
        parts = [*captions, item.get("table_body") or "",
                 *(item.get("table_footnote") or [])]
        block = "\n\n".join(p.strip() for p in parts if p and p.strip())
        return (block, block) if block else None
    if kind == "image":
        block = f"[Gambar] {item.get('img_path', '')}\n{item.get('img_description', '')}".strip()
        return block, block
    return None


def render_pages(content_list: list, raw_parse: dict, n_pages: int) -> tuple[list, list]:
    """One markdown and one plain string per page, ``""`` for a page with nothing."""
    md = [[] for _ in range(n_pages)]
    plain = [[] for _ in range(n_pages)]
    recent = [[] for _ in range(n_pages)]   # the last text blocks of each page
    for item in content_list:
        page = item.get("page_idx")
        if not isinstance(page, int) or not 0 <= page < n_pages:
            continue
        block = _item_block(item, _norm(" ".join(recent[page][-3:])))
        if block is None:
            continue
        md[page].append(block[0])
        plain[page].append(block[1])
        recent[page] = recent[page] + [block[1]] if item.get("type") == "text" else []
    raw_pages = {p.get("page_idx"): p for p in (raw_parse or {}).get("pages", [])}
    markdown = []
    for page in range(n_pages):
        raw_page = raw_pages.get(page, {})
        parts = (_page_layout_texts(raw_page, HEADER_TYPES) + md[page]
                 + _page_layout_texts(raw_page, FOOTER_TYPES))
        markdown.append("\n\n".join(parts))
    return markdown, ["\n".join(p) for p in plain]


def heading_items(content_list: list) -> list:
    """Textract headings as the flat items ``utils.post_processing`` turns into a tree.

    TITLE (level 1) is a top-level node, SECTION_HEADER (level 2) sits under the
    latest TITLE (top-level before the first one). A heading repeated on two or
    more pages is a letterhead and is dropped.
    """
    first_on_page = {}
    for index, item in enumerate(content_list):
        if _item_block(item) is not None:
            first_on_page.setdefault(item.get("page_idx"), index)

    candidates = [
        (index, item) for index, item in enumerate(content_list)
        if item.get("type") == "text" and item.get("text_level") in (1, 2)
        and HEADING_MIN_CHARS <= len((item.get("text") or "").strip()) <= HEADING_MAX_CHARS
    ]
    pages_of = {}
    for _, item in candidates:
        pages_of.setdefault(item["text"].strip().casefold(), set()).add(item.get("page_idx"))

    items, top, sub, under_title = [], 0, 0, False
    for index, item in candidates:
        title = item["text"].strip()
        if len(pages_of[title.casefold()]) > 1:
            continue
        if item["text_level"] == 1 or not under_title:
            top, sub = top + 1, 0
            under_title = under_title or item["text_level"] == 1
            code = str(top)
        else:
            sub += 1
            code = f"{top}.{sub}"
        items.append({
            "structure": code,
            "title": title,
            "physical_index": item["page_idx"] + 1,
            "appear_start": "yes" if first_on_page.get(item["page_idx"]) == index else "no",
        })
    return items


def build_structure(content_list: list, n_pages: int) -> list:
    """Tree ``{title, start_index, end_index, nodes}`` from the Textract headings."""
    from pageindex.utils import cover_subtree_ranges, post_processing
    items = heading_items(content_list)
    if not items:
        return []
    return cover_subtree_ranges(post_processing(items, n_pages))


# ── spreadsheets ──

def remap_sheets(content_list: list, sheet_names: list[str]) -> None:
    """Point each item's ``page_idx`` (the parser's sheet name) at the sheet's 0-based position."""
    position = {name: index for index, name in enumerate(sheet_names)}
    for item in content_list:
        item["page_idx"] = position[item["page_idx"]]


def sheet_structure(sheet_names: list[str]) -> list:
    """One tree node per sheet, titled by its name: a sheet is a page."""
    return [{"title": name, "start_index": page, "end_index": page}
            for page, name in enumerate(sheet_names, 1)]


# ── the extract_toc hook ──

_original_extract_toc = None


def _extract_toc(doc_handle, *args, **kwargs):
    """``pageindex.flash.main.extract_toc``, with the Textract parse of the document
    being indexed (``_CURRENT``) folded in; unchanged for every other document."""
    parsed = _CURRENT.get()
    if parsed is not None and parsed.structure is not None:
        # a spreadsheet: its PDF is blank, the sheets are the tree
        return {"doc_name": Path(str(doc_handle)).name, "doc_title": None,
                "structure": copy.deepcopy(parsed.structure),
                "has_abstract_or_references_section": False,
                "page_texts": list(parsed.plain), "toc_source": "sheets"}
    result = _original_extract_toc(doc_handle, *args, **kwargs)
    if parsed is None:
        return result
    texts = result.get("page_texts")
    if not texts:
        # bookmark-only extraction: no page text at all
        result["page_texts"] = list(parsed.plain)
    elif len(texts) == len(parsed.plain):
        # pdfium's text stays on the other pages: Flash's titles were matched against it
        for page in parsed.visual_pages:
            texts[page] = parsed.plain[page]
    else:
        logger.warning(f"[multimodal] {result.get('doc_name')}: Flash sees {len(texts)} pages, "
                       f"Textract {len(parsed.plain)}; summaries use Flash's text")
    if not result.get("structure"):
        # no text layer and no bookmarks; empty here, Flash falls back to a node per page
        result["structure"] = build_structure(parsed.content_list, len(parsed.plain))
    return result


_extract_toc.multimodal = True


def install() -> None:
    """Route Flash's extraction through ``_extract_toc``; idempotent."""
    global _original_extract_toc
    import pageindex.flash.api as flash_api
    if getattr(flash_api.extract_toc, "multimodal", False):
        return
    _original_extract_toc = flash_api.extract_toc
    flash_api.extract_toc = _extract_toc


# ── the API ──

class MultimodalLocalAPI(LocalAPI):
    """Made by ``MultimodalClient`` swapping the class of its ``LocalAPI``, so the
    constructor stays upstream's; ``_init_multimodal`` adds this class's state."""

    def _init_multimodal(self) -> None:
        self._lock = threading.Lock()
        # keyed by the work-dir PDF a submit indexes (str path, name for _display_names)
        self._sources: dict[str, Path] = {}      # -> the submitted file
        self._parsed: dict[str, Parsed] = {}     # -> its Textract / spreadsheet parse
        self._display_names: dict[str, str] = {}  # -> the submitted file's name

    def submit_document(self, file_path: str, mode: str | None = None, **kwargs):
        src = Path(os.path.abspath(os.path.expanduser(str(file_path))))
        ext = src.suffix.lower()
        if not src.is_file() or not (ext == ".pdf" or ext in IMAGE_EXTS
                                     or ext in SPREADSHEET_EXTS or _is_office_file(src)):
            # upstream's own error: no such file, or not a PDF
            return super().submit_document(file_path, mode=mode, **kwargs)
        tmp = Path(tempfile.mkdtemp(prefix="pageindex_mm_"))
        # a fresh name per submit: two submits of one file never share a work PDF
        pdf = tmp / f"{uuid.uuid4().hex}.pdf"
        try:
            if ext in IMAGE_EXTS:
                image_to_pdf(src, pdf)
            elif ext in SPREADSHEET_EXTS:
                _blank_pdf(pdf, len(_sheet_names(src)))
            elif ext == ".pdf":
                shutil.copyfile(src, pdf)
            else:
                _office_to_pdf(src, pdf)
            with self._lock:
                self._sources[str(pdf)] = src
                self._display_names[pdf.name] = sanitize_filename(src.name)
            return super().submit_document(str(pdf), mode=mode, **kwargs)
        finally:
            with self._lock:
                self._sources.pop(str(pdf), None)
                self._parsed.pop(str(pdf), None)
                self._display_names.pop(pdf.name, None)
            _discard_workdir(tmp, src.stem)

    def _unique_doc_name(self, name: str) -> str:
        with self._lock:
            name = self._display_names.get(name, name)
        return super()._unique_doc_name(name)

    def _extract_page_texts(self, file_path: str) -> list[str]:
        with self._lock:
            src = self._sources.get(file_path)
        if src is None:
            return LocalAPI._extract_page_texts(file_path)
        pdf, tmp = Path(file_path), Path(file_path).parent
        if src.suffix.lower() in SPREADSHEET_EXTS:
            parsed = run_off_loop(self._parse_spreadsheet, src, tmp)
        else:
            pages = textract_pages(pdf)
            image = src if src.suffix.lower() in IMAGE_EXTS else None
            parsed = run_off_loop(self._parse, pdf, tmp, pages, image) if pages else None
            if parsed is None:
                return LocalAPI._extract_page_texts(file_path)
        with self._lock:
            self._parsed[file_path] = parsed
        return list(parsed.markdown)

    def _index_flash(self, file_path: str):
        # runs on the thread page_index_flash (and so the hook) runs on
        with self._lock:
            parsed = self._parsed.get(file_path)
        token = _CURRENT.set(parsed)
        try:
            return super()._index_flash(file_path)
        finally:
            _CURRENT.reset(token)

    # ── parsers ──

    def _parse(self, pdf: Path, tmp: Path, pages: list[int],
               image: Path | None = None) -> Parsed | None:
        """Textract on ``pages`` → content list → S3 → VLM → pages; None keeps the original text.

        Only ``pages`` go to Textract (as a sub-PDF); every other page keeps the
        original PyPDF2 text, byte for byte. ``image`` is the submitted image file
        when ``pdf`` was made from one: a visual image goes to the VLM whole.
        """
        import pymupdf
        from .parser.figures import host_figures
        from .parser.pdf.amazon_textract import parse_document_to_json, reconstruct_content_list
        from .vlm import describe_figures
        try:
            original = LocalAPI._extract_page_texts(str(pdf))
        except Exception:
            return None   # the original path reports an unreadable PDF its own way
        with pymupdf.open(pdf) as doc:
            n_pages = doc.page_count
        if len(original) != n_pages:
            logger.warning(f"[multimodal] {pdf.name}: PyPDF2 sees {len(original)} pages, "
                           f"PyMuPDF {n_pages}; pakai teks asli (PyPDF2)")
            return None
        try:
            source = pdf
            if len(pages) < n_pages:
                source = tmp / "textract_pages.pdf"
                with pymupdf.open(pdf) as doc:
                    doc.select(pages)
                    doc.save(source)
            raw = asyncio.run(parse_document_to_json(
                str(source), output_file=str(tmp / "raw_parsing_result.json")))
            content_list = reconstruct_content_list(
                raw, output_file=str(tmp / "formatted_parsing_result.json"))
            remap_pages(content_list, raw, pages)
            if image is not None and text_coverage(raw["pages"][0]) < VISUAL_TEXT_COVERAGE:
                content_list = as_whole_image(content_list, image_for_vlm(image, tmp))
            host_figures(content_list)
            asyncio.run(describe_figures(content_list))
        except Exception as e:
            if has_text_layer(pdf):
                logger.warning(f"[multimodal] {pdf.name}: Textract/S3/VLM gagal, "
                               f"pakai teks asli (PyPDF2): {e}", exc_info=True)
                return None
            raise PageIndexAPIError(
                f"Failed to submit document: OCR (Textract/VLM) gagal untuk PDF "
                f"tanpa text layer: {e}") from e
        markdown, plain = render_pages(content_list, raw, n_pages)
        return Parsed(content_list, merge_pages(markdown, original, pages),
                      merge_pages(plain, original, pages), visual_pages=list(pages))

    def _parse_spreadsheet(self, src: Path, tmp: Path) -> Parsed:
        """Parser → sheets as pages → S3 → VLM. Fails rather than indexing a workbook
        without its figures: a spreadsheet has no other text to fall back on."""
        from .parser.figures import host_figures
        from .parser.spreadsheet.parser import parse_xlsx
        from .vlm import describe_figures
        try:
            sheet_names = _sheet_names(src)
            content_list = parse_xlsx(src, tmp / "images")
            remap_sheets(content_list, sheet_names)
            host_figures(content_list)
            asyncio.run(describe_figures(content_list))
        except Exception as e:
            raise PageIndexAPIError(f"Failed to submit document: could not read spreadsheet: {e}") from e
        markdown, plain = render_pages(content_list, {}, len(sheet_names))
        if not any(text.strip() for text in plain):
            raise PageIndexAPIError(
                "Failed to submit document: spreadsheet has no content. All sheets are blank.")
        markdown = [f"{SHEET_PREFIX}{name}\n\n{text}" for name, text in zip(sheet_names, markdown)]
        return Parsed(content_list, markdown, plain, structure=sheet_structure(sheet_names))


# ── submit helpers ──

def _sheet_names(src: Path) -> list[str]:
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(src, read_only=True)
    except Exception as e:
        raise PageIndexAPIError(f"Failed to submit document: could not read spreadsheet: {e}") from e
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _blank_pdf(dst: Path, n_pages: int) -> Path:
    """A PDF of ``n_pages`` empty pages: a spreadsheet's stand-in, a page per sheet."""
    import pymupdf
    with pymupdf.open() as doc:
        for _ in range(n_pages):
            doc.new_page()
        doc.save(dst)
    return dst


def _is_office_file(path: Path) -> bool:
    from .parser.office.office_to_pdf import SUPPORTED_EXTENSIONS
    return path.suffix.lstrip(".").lower() in SUPPORTED_EXTENSIONS


def _office_to_pdf(src: Path, dst: Path) -> Path:
    """The PDF LibreOffice (unoserver) renders from an office file, moved to ``dst``."""
    from .parser.office.office_to_pdf import convert_office_to_pdf
    try:
        rendered = Path(run_off_loop(asyncio.run, convert_office_to_pdf(str(src))))
    except Exception as e:
        raise PageIndexAPIError(
            f"Failed to submit document: could not convert {src.name} to PDF: {e}") from e
    shutil.move(rendered, dst)
    shutil.rmtree(rendered.parent, ignore_errors=True)
    return dst


def _discard_workdir(tmp: Path, stem: str) -> None:
    """Remove a submit's work dir, keeping a copy under ``PAGEINDEX_MM_DEBUG_DIR`` when set."""
    debug_dir = os.environ.get(DEBUG_DIR_ENV)
    if debug_dir:
        shutil.copytree(tmp, Path(debug_dir) / stem, dirs_exist_ok=True)
    shutil.rmtree(tmp, ignore_errors=True)
