import datetime as dt
import io
import os
import re
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from unstructured.partition.xlsx import partition_xlsx
from unstructured.documents.elements import Table
from html_table_rescuer import TableParser
from bs4 import BeautifulSoup

from ..figures import is_large_image, save_for_vlm



# ============================================================
# 0. PARTISI PER SHEET (via file temporary), nilai seperti tampil di Excel
# ============================================================
_PERCENT_DECIMALS = re.compile(r"0\.(0+)%")
_TIME_TOKENS = re.compile(r"[hs]", re.IGNORECASE)


def display_value(value, number_format: str) -> str | None:
    """Teks yang Excel tampilkan untuk sel persen atau tanggal; None untuk sel lain.

    Persen yang tersimpan 0.333333 dengan format ``0.0%`` dibaca "33.3%", bukan
    0.333333. Tanggal kehilangan jam 00:00:00 yang tidak pernah tampil, kecuali
    formatnya memang menampilkan jam. Angka lain dibiarkan mentah: "4.680" atau
    "4,680" terbaca sebagai angka berbeda tergantung locale pembacanya.
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and "%" in number_format:
        match = _PERCENT_DECIMALS.search(number_format)
        places = Decimal(1).scaleb(-len(match.group(1))) if match else Decimal(1)
        percent = (Decimal(str(value)) * 100).quantize(places, rounding=ROUND_HALF_UP)
        return f"{percent}%"
    if isinstance(value, dt.datetime) and not _TIME_TOKENS.search(number_format):
        return value.date().isoformat()
    return None


def show_as_displayed(ws) -> None:
    """Tulis sel persen dan tanggal sebagai teks, seperti yang tampil di Excel."""
    for row in ws.iter_rows():
        for cell in row:
            text = display_value(cell.value, cell.number_format)
            if text is not None:
                cell.value = text


def partition_single_sheet(filepath: str, sheet_name: str):
    """Partisi HANYA satu sheet dengan membuat file XLSX sementara."""
    wb = load_workbook(filepath, data_only=True)
    for s in list(wb.sheetnames):
        if s != sheet_name:
            del wb[s]
    show_as_displayed(wb[sheet_name])  # hanya salinan sementara; file asli tidak berubah

    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(tmp_fd)
    try:
        wb.save(tmp_path)
        elements = partition_xlsx(
            filename=tmp_path,
            infer_table_structure=True,
            find_subtable=True,
        )
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)

    return elements


# ============================================================
# 0a. PROMOSI BARIS PERTAMA MENJADI <th>
# ============================================================
def promote_first_row_to_header(html: str) -> str:
    """Ubah baris pertama tabel dari <td> menjadi <th> dan bungkus <thead>."""
    if not html:
        return html
    try:
        soup = BeautifulSoup(html, "lxml")
        table = soup.find("table")
        if not table:
            return html

        if table.find("thead") or table.find("th"):
            return html

        rows = table.find_all("tr")
        if not rows:
            return html

        first_row = rows[0]
        cells = first_row.find_all("td")
        if not cells:
            return html

        thead = soup.new_tag("thead")
        new_tr = soup.new_tag("tr")
        for cell in cells:
            th = soup.new_tag("th")
            th.string = cell.get_text(strip=True)
            new_tr.append(th)
        thead.append(new_tr)

        first_row.extract()
        table.insert(0, thead)

        return str(table)
    except Exception:
        return html


# ============================================================
# 0b. KONVERSI HTML TABEL -> MARKDOWN
# ============================================================
def html_table_to_markdown(html: str) -> str | None:
    """Konversi HTML tabel dari unstructured.metadata.text_as_html menjadi Markdown."""
    if not html:
        return None
    try:
        html_with_header = promote_first_row_to_header(html)
        parser = TableParser(html_with_header)
        tables = parser.parse()
        if not tables:
            return None
        md_tables = []
        for t in tables:
            try:
                md_tables.append(t.to_markdown())
            except Exception:
                continue
        if not md_tables:
            return None
        return "\n\n".join(md_tables)
    except Exception:
        return None


# ============================================================
# 1. EKSTRAKSI GAMBAR & HYPERLINK DENGAN POSISI BARIS
# ============================================================
def extract_sheet_assets(filepath, images_dir):
    """Return {sheet_name: {"images": [...], "hyperlinks": [...]}} dengan row/col.

    Gambar di bawah ambang ``FIGURE_MIN_*`` dilewati: tidak ditulis ke ``images_dir``.
    """
    wb = load_workbook(filepath, data_only=False)
    img_dir = Path(images_dir)
    img_dir.mkdir(parents=True, exist_ok=True)
    result = {}

    for ws in wb.worksheets:
        images = []
        hyperlinks = []

        # --- Gambar ---
        for img in ws._images:
            try:
                row = img.anchor._from.row + 1
                col = img.anchor._from.col + 1
                img_data = img._data()
                ext = "png"
                if getattr(img, "format", None):
                    ext = img.format.lower()
                img_filename = f"{ws.title.replace(' ', '_')}_r{row}c{col}.{ext}"
                if not is_large_image(img_data, img_filename):
                    continue
                img_path = img_dir / img_filename
                with Image.open(io.BytesIO(img_data)) as pil_img:
                    img_path = save_for_vlm(pil_img, img_path, original=img_data)

                # Cari label konteks: nilai terdekat di kiri, lalu di atas
                label = None
                for c in range(col, 0, -1):
                    v = ws.cell(row=row, column=c).value
                    if v and str(v).strip():
                        label = str(v).strip()
                        break
                if not label:
                    for r in range(row, 0, -1):
                        v = ws.cell(row=r, column=col).value
                        if v and str(v).strip():
                            label = str(v).strip()
                            break

                images.append({
                    "row": row,
                    "col": col,
                    "cell": f"{get_column_letter(col)}{row}",
                    "label": label,
                    "path": str(img_path),
                })
            except Exception as e:
                images.append({"error": str(e), "row": 0, "col": 0})

        # --- Hyperlink ---
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink:
                    hyperlinks.append({
                        "row": cell.row,
                        "col": cell.column,
                        "cell": cell.coordinate,
                        "text": str(cell.value) if cell.value else "",
                        "target": cell.hyperlink.target,
                        "tooltip": cell.hyperlink.tooltip,
                    })

        images.sort(key=lambda x: (x.get("row", 0), x.get("col", 0)))
        hyperlinks.sort(key=lambda x: (x["row"], x["col"]))
        result[ws.title] = {"images": images, "hyperlinks": hyperlinks}

    return result


# ============================================================
# 2. CARI BARIS AWAL & AKHIR ELEMEN DI SHEET
# ============================================================
def element_start_row(ws, element_text, start_row=1):
    """Cari baris awal elemen berdasarkan token unik dari teksnya."""
    if not element_text:
        return start_row
    snippet = element_text.strip()[:80]
    tokens = [t for t in snippet.split() if len(t) > 2][:3]
    if not tokens:
        return start_row

    for r in range(start_row, ws.max_row + 1):
        row_text = " ".join(
            str(ws.cell(row=r, column=c).value or "")
            for c in range(1, ws.max_column + 1)
        )
        if all(tok in row_text for tok in tokens):
            return r
    return start_row


def element_end_row(el, start_row):
    """Untuk Table, hitung jumlah <tr>; untuk lainnya, 1 baris."""
    if not isinstance(el, Table):
        return start_row
    html = getattr(el.metadata, "text_as_html", "") or ""
    tr_count = html.count("<tr>")
    return start_row + max(tr_count - 1, 0)


# ============================================================
# 3. PARSE
# ============================================================
def parse_xlsx(path: str | Path, images_dir: str | Path) -> list[dict]:
    """A .xlsx file → content list (text / table / image), sheet by sheet.

    ``page_idx`` is the sheet name. Embedded images are written to ``images_dir`` and
    returned as ``image`` items with a local ``img_path``; images under the
    ``FIGURE_MIN_*`` thresholds are skipped. Hosting them (S3) is the caller's job.
    """
    path = str(path)
    assets = extract_sheet_assets(path, images_dir)

    wb = load_workbook(path, data_only=True)
    sheet_names = wb.sheetnames

    output = []

    for sheet_name in sheet_names:
        elements = partition_single_sheet(path, sheet_name)

        ws = wb[sheet_name]
        sheet_assets = assets.get(sheet_name, {"images": [], "hyperlinks": []})

        # Gabung semua aset, urutkan berdasarkan (row, col)
        all_assets = []
        for img in sheet_assets["images"]:
            if "error" in img:
                continue
            all_assets.append({
                "type": "image",
                "row": img["row"],
                "col": img["col"],
                "cell": img["cell"],
                "path": str(Path(img["path"]).resolve()),
                "label": img.get("label"),
            })
        for link in sheet_assets["hyperlinks"]:
            all_assets.append({
                "type": "link",
                "row": link["row"],
                "col": link["col"],
                "cell": link["cell"],
                "target": link["target"],
                "text": link["text"],
                "tooltip": link.get("tooltip"),
            })
        all_assets.sort(key=lambda x: (x["row"], x["col"]))

        # DEDUP: URL yang sudah dirender sebagai aset
        asset_urls = set()
        for link in sheet_assets["hyperlinks"]:
            if link.get("target"):
                asset_urls.add(link["target"].strip())
            if link.get("text"):
                asset_urls.add(link["text"].strip())

        def is_asset_url(text: str) -> bool:
            if not text:
                return False
            return text.strip() in asset_urls

        # Tentukan rentang baris setiap elemen
        element_spans = []
        cursor = 1
        for el in elements:
            start = element_start_row(ws, el.text, cursor)
            end = element_end_row(el, start)
            element_spans.append((start, end))
            cursor = end + 1

        emitted = set()

        def emit_assets_until(max_row):
            """Emit aset yang row-nya <= max_row dan belum pernah di-emit."""
            for a in all_assets:
                key = (a["row"], a["col"], a["type"])
                if key in emitted:
                    continue
                if a["row"] <= max_row:
                    if a["type"] == "image":
                        output.append({
                            "type": "image",
                            "img_path": a["path"],
                            "image_caption": [],
                            "image_footnote": [],
                            "page_idx": sheet_name,
                        })
                    elif a["type"] == "link":
                        # ✅ Link ditulis sebagai text block
                        link_text = (a["text"] or "").strip()
                        target = (a["target"] or "").strip()
                        if link_text and target and link_text != target:
                            text_value = f"{link_text} ({target})"
                        else:
                            text_value = target or link_text
                        if text_value:
                            output.append({
                                "type": "text",
                                "text": text_value,
                                "page_idx": sheet_name,
                            })
                    emitted.add(key)

        for idx, el in enumerate(elements):
            start_row, end_row = element_spans[idx]
            emit_assets_until(end_row)

            # SKIP: elemen non-Table yang teksnya adalah URL aset
            if not isinstance(el, Table) and is_asset_url(el.text):
                continue

            if isinstance(el, Table):
                html = getattr(el.metadata, "text_as_html", None)
                md_table = html_table_to_markdown(html) if html else None

                table_body = md_table if md_table else (html or "")
                output.append({
                    "type": "table",
                    "table_body": table_body,
                    "table_caption": [],
                    "table_footnote": [],
                    "page_idx": sheet_name,
                })
            else:
                text = (el.text or "").strip()
                if text:
                    output.append({
                        "type": "text",
                        "text": text,
                        "page_idx": sheet_name,
                    })

        # Sisa aset yang row-nya setelah elemen terakhir
        emit_assets_until(10**9)

    return output